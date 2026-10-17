"""Measure every site in heatwise_pakistan_sites.xlsx with HeatWise's real data layer.

Put this file in the backend folder, then (one-time):  pip install openpyxl
Examples:
    py probe_from_excel.py
    py probe_from_excel.py heatwise_pakistan_sites.xlsx
    py probe_from_excel.py heatwise_pakistan_sites.xlsx --city Lahore
    py probe_from_excel.py heatwise_pakistan_sites.xlsx --date 2026-06-10 --save results.csv

It needs internet. Each site takes a few seconds (public OpenStreetMap server), and
results are cached, so a second run is fast. Read the 'adj' column: you want sites in the
same city whose adjustment differs by ~2 C or more, with conf = high or medium.
"""
from __future__ import annotations

import argparse
import csv
import datetime
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook  # noqa: E402

from services import open_data_service as od  # noqa: E402


def read_sites(path: str, city: str | None):
    ws = load_workbook(path, data_only=True)["Sites"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        name, c, prov, lat, lon = r[0], r[1], r[2], r[3], r[4]
        if not name or not isinstance(lat, (int, float)) or not isinstance(lon, (int, float)):
            continue
        if city and (c or "").lower() != city.lower():
            continue
        rows.append((name, c, float(lat), float(lon)))
    return rows


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", default="heatwise_pakistan_sites.xlsx",
                    help="path to the workbook (default: heatwise_pakistan_sites.xlsx in this folder)")
    ap.add_argument("--city", help="only sites in this city, e.g. Lahore")
    ap.add_argument("--date", default=datetime.date.today().isoformat(), help="YYYY-MM-DD (default today)")
    ap.add_argument("--hour", default="14:00")
    ap.add_argument("--save", help="also write results to this CSV file")
    a = ap.parse_args()

    day = datetime.date.fromisoformat(a.date)
    sites = read_sites(a.xlsx, a.city)
    if not sites:
        print("No sites found (check the file path and --city spelling).")
        return 1

    print(f"Probing {len(sites)} site(s) for {day} at {a.hour}\n")
    print(f"{'Site':34}{'City':12}{'paved%':>8}{'veg%':>7}{'adj C':>8}{'temp C':>8}  conf")
    out = []
    for name, city, lat, lon in sites:
        try:
            s = od.get_site_snapshot(lat, lon, day, a.hour)
            l = s["land"]
            print(f"{name[:33]:34}{(city or '')[:11]:12}{l['impervious_pct']:8.1f}{l['tree_canopy_pct']:7.1f}"
                  f"{s['adjustment_c']:+8.1f}{s['temperature_c']:8.1f}  {l['confidence']}")
            out.append([name, city, lat, lon, round(l["impervious_pct"], 1), round(l["tree_canopy_pct"], 1),
                        s["adjustment_c"], round(s["temperature_c"], 1), l["confidence"]])
        except Exception as exc:  # noqa: BLE001
            print(f"{name[:33]:34}{(city or '')[:11]:12}  FAILED: {exc}")
            out.append([name, city, lat, lon, "", "", "", "", f"failed: {exc}"])
        time.sleep(1.0)   # be polite to the free public map server

    if a.save:
        with open(a.save, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["name", "city", "lat", "lon", "paved_pct", "vegetation_pct", "adjustment_c", "temp_c", "confidence"])
            w.writerows(out)
        print(f"\nSaved {a.save}")
    print("\nPick sites in one city whose 'adj C' differs by about 2 or more, with conf high or medium.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())